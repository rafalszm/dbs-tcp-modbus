from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CSV_COLUMNS = [
    "key",
    "name",
    "function",
    "address",
    "type",
    "bit",
    "unit",
    "scale",
    "offset",
    "precision",
    "device_class",
    "state_class",
    "icon",
    "section",
    "enabled_by_default",
    "count",
    "word_order",
    "byte_order",
]

STATIONS = [
    {
        "sheet_contains": "Czarn",
        "prefix": "cz",
        "output": Path("examples/czarnow.csv"),
    },
    {
        "sheet_contains": "Gawartowa",
        "prefix": "gw",
        "output": Path("examples/gawartowa_wola.csv"),
    },
]

TYPE_MAP = {
    "bool": "bit",
    "real": "float32",
    "udint": "uint32",
    "dint": "int32",
    "int": "int16",
    "uint": "uint16",
    "word": "uint16",
}


def main() -> int:
    parser = argparse.ArgumentParser(description="Build DBS TCP Modbus CSV maps from the automation workbook.")
    parser.add_argument("workbook", type=Path, help="Path to MB TCP Leszno.xlsx")
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, data_only=False)
    summaries: list[dict[str, Any]] = []

    for station in STATIONS:
        sheet = find_sheet(workbook, station["sheet_contains"])
        rows, summary = build_rows(sheet, station["prefix"])
        output_path = args.repo_root / station["output"]
        write_csv(output_path, rows)
        summary["output"] = str(station["output"]).replace("\\", "/")
        summaries.append(summary)

    for summary in summaries:
        print(f"{summary['sheet']}: {summary['output']}")
        print(
            f"  written={summary['written']} skipped_static_or_reserve={summary['skipped_static']} "
            f"bit_rows={summary['types'].get('bit', 0)} 32bit_low_high={summary['low_high']}"
        )
        if summary["overlaps"]:
            print("  source overlaps:")
            for overlap in summary["overlaps"]:
                print(f"    address {overlap['address']}: {', '.join(overlap['kinds'])}")
        if summary["ip"] or summary["port"]:
            print(f"  workbook endpoint: {summary['ip'] or '?'}:{summary['port'] or '?'}")

    return 0


def find_sheet(workbook: Any, needle: str) -> Any:
    for sheet in workbook.worksheets:
        if needle.lower() in slugify(sheet.title).replace("_", ""):
            return sheet
    for sheet in workbook.worksheets:
        if needle.lower() in sheet.title.lower():
            return sheet
    raise ValueError(f"Workbook does not contain a sheet matching {needle!r}")


def build_rows(sheet: Any, prefix: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    keys: set[str] = set()
    skipped_static = 0
    type_counter: Counter[str] = Counter()
    address_kinds: dict[int, set[str]] = defaultdict(set)

    current_address: int | None = None
    current_section = ""
    bool_index = 0

    for row_number in range(1, sheet.max_row + 1):
        marker = cell_text(sheet, row_number, 1)
        name = cell_text(sheet, row_number, 2)
        source_type = cell_text(sheet, row_number, 3).lower()
        address = parse_int(cell_text(sheet, row_number, 4))
        section_marker = cell_text(sheet, row_number, 5)
        note = cell_text(sheet, row_number, 6)
        unit_hint = cell_text(sheet, row_number, 7)
        extra_note = cell_text(sheet, row_number, 8)

        value_type = TYPE_MAP.get(source_type)
        if value_type is None and section_marker and address is None:
            current_section = clean_section(section_marker)
            continue

        is_header = address is not None and value_type is None
        if is_header:
            current_address = address
            bool_index = 0
            current_section = clean_section(section_marker or note or name)
            continue

        if not name or value_type is None:
            continue

        if value_type == "bit":
            if address is not None:
                current_address = address
                bool_index = 0
                row_section = clean_section(section_marker)
                if not row_section and looks_like_section_note(note):
                    row_section = clean_section(note)
                if row_section:
                    current_section = row_section
            if current_address is None:
                raise ValueError(f"{sheet.title} row {row_number}: Bool row without current address")
            bit = 15 - bool_index
            bool_index += 1
            row_address = current_address
        else:
            if address is None:
                raise ValueError(f"{sheet.title} row {row_number}: {name} has no address")
            bit = ""
            row_address = address

        if is_static_or_reserve(name, note):
            skipped_static += 1
            continue

        if value_type == "bit":
            unit = ""
            device_class = ""
            state_class = ""
        else:
            unit = infer_unit(name, note, unit_hint, extra_note)
            device_class = infer_device_class(name, unit)
            state_class = infer_state_class(value_type, name)
        section = current_section if value_type == "bit" else clean_section(section_marker)
        key = unique_key(prefix, section, name, row_address, keys)
        csv_row = {
            "key": key,
            "name": name.strip(),
            "function": "3",
            "address": str(row_address),
            "type": value_type,
            "bit": str(bit),
            "unit": unit,
            "scale": "1",
            "offset": "0",
            "precision": "3" if value_type == "float32" else "",
            "device_class": device_class,
            "state_class": state_class,
            "icon": "",
            "section": section,
            "enabled_by_default": "true",
            "count": "",
            "word_order": "low_high" if value_type in {"float32", "uint32", "int32"} else "",
            "byte_order": "big" if value_type in {"float32", "uint32", "int32"} else "",
        }
        rows.append(csv_row)
        type_counter[value_type] += 1
        address_kinds[row_address].add(value_type)

    overlaps = [
        {"address": address, "kinds": sorted(kinds)}
        for address, kinds in sorted(address_kinds.items())
        if "bit" in kinds and any(kind != "bit" for kind in kinds)
    ]
    summary = {
        "sheet": sheet.title,
        "ip": cell_text(sheet, 2, 2),
        "port": cell_text(sheet, 2, 3),
        "written": len(rows),
        "skipped_static": skipped_static,
        "types": dict(type_counter),
        "low_high": sum(1 for row in rows if row["word_order"] == "low_high"),
        "overlaps": overlaps,
    }
    return rows, summary


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def cell_text(sheet: Any, row: int, column: int) -> str:
    value = sheet.cell(row, column).value
    return "" if value is None else str(value).strip()


def parse_int(value: str) -> int | None:
    if not value:
        return None
    try:
        return int(float(value))
    except ValueError:
        return None


def clean_section(value: str) -> str:
    value = re.sub(r"\bcd\b", "", value or "", flags=re.IGNORECASE).strip(" -_/")
    if value.lower() in {"wejscia sterownika", "wejścia sterownika", "rezerwa"}:
        return ""
    return value


def looks_like_section_note(value: str) -> bool:
    text = (value or "").strip().lower()
    return bool(text) and not text.startswith("1-") and text not in {"rezerwa"}


def is_static_or_reserve(name: str, note: str) -> bool:
    name_slug = slugify(name)
    return name_slug.startswith("static") or (note or "").strip().lower() == "rezerwa"


def infer_unit(name: str, note: str, unit_hint: str, extra_note: str) -> str:
    text = " ".join(part for part in [name, note, unit_hint, extra_note] if part).lower()
    name_lower = name.lower()

    if name_lower == "kwh":
        return "kWh"
    if name_lower == "kvah":
        return "kVAh"
    if name_lower == "kvarh":
        return "kvarh"
    if "m3/h" in text:
        return "m3/h"
    if re.search(r"\bm[345678]\b", text) or "m3" in text:
        return "m3"
    if "bar" in text or "cis_" in name_lower:
        return "bar"
    if "%" in text or "poz_przep" in name_lower:
        return "%"
    if re.search(r"\bm\b", text) or name_lower.startswith("poz_") or name_lower.startswith("poziom_"):
        return "m"
    if "sekund" in text or "do_konca_s" in name_lower:
        return "s"
    if (
        "hh:mm" not in text
        and (
            unit_hint.lower() == "min"
            or extra_note.lower() == "min"
            or " min" in f" {note.lower()} "
            or name_lower.startswith("czas_pr_min")
        )
    ):
        return "min"
    if "godzin" in text or "_godz" in name_lower:
        return "h"
    if name_lower.startswith("voltage_") or name_lower.startswith("avg_voltage"):
        return "V"
    if name_lower.startswith("current_") or name_lower.startswith("avg_current"):
        return "A"
    if name_lower == "frequency":
        return "Hz"
    if name_lower.startswith("kw") or name_lower == "total_kw":
        return "kW"
    if name_lower.startswith("kva") or name_lower == "total_kva":
        return "kVA"
    if name_lower.startswith("kvar") or name_lower == "total_kvar":
        return "kvar"
    return ""


def infer_device_class(name: str, unit: str) -> str:
    name_lower = name.lower()
    if unit == "V":
        return "voltage"
    if unit == "A":
        return "current"
    if unit == "Hz":
        return "frequency"
    if unit == "bar":
        return "pressure"
    if unit == "m":
        return "distance"
    if unit == "kW":
        return "power"
    if unit == "kWh":
        return "energy"
    if "temperature" in name_lower or "temp" in name_lower:
        return "temperature"
    return ""


def infer_state_class(value_type: str, name: str) -> str:
    if value_type == "bit":
        return ""
    name_lower = name.lower()
    if name_lower in {"kwh", "kvah", "kvarh"} or name_lower.startswith("licz"):
        return "total_increasing"
    return "measurement"


def unique_key(prefix: str, section: str, name: str, address: int, keys: set[str]) -> str:
    base = slugify("_".join(part for part in [prefix, section, name, str(address)] if part))
    key = base
    suffix = 2
    while key in keys:
        key = f"{base}_{suffix}"
        suffix += 1
    keys.add(key)
    return key


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-zA-Z0-9_]+", "_", normalized).strip("_").lower()
    return re.sub(r"_+", "_", slug) or "value"


if __name__ == "__main__":
    sys.exit(main())
